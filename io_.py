import tkinter as tk
from tkinter import ttk,Tk
import pathlib
import pygubu
from app_config import (
    APP_SOFTWARE_MODEL_REV,
    DATA_DIR_NAME,
    STARTUP_JSON_NAME,
    USB_TRANSFER_CONNECTED_DEFAULT,
    USB_TRANSFER_ENABLED_DEFAULT,
    WIFI_TRANSFER_CONNECTED_DEFAULT,
    WIFI_TRANSFER_ENABLED_DEFAULT,
)
from conversions import CONVERSIONS
import time
import datetime
import openpyxl 
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
import json
import platform 
import os

PROJECT_PATH = pathlib.Path(__file__).parent

class I_O:
    
    # False = Prox Left unchecked (distal-first Excel/tree headers).
    cone_flip = False
    Software_Model_Rev = APP_SOFTWARE_MODEL_REV
     
    if (platform.system()=="Windows"): 
        path_1 = str(DATA_DIR_NAME + '\\')
    else:
        path_1 = str(DATA_DIR_NAME + '/')
  #  path_1 = "/RC_Measurment_Sys_9_20_23/I_O/"

    def __init__(self, parent):
        self.parent = parent

    @staticmethod
    def cone_flip_from_excel_headers(header_row):
        """Infer Prox Left from Excel header row (column 2 / index 1)."""
        if not header_row or len(header_row) < 2:
            return False
        col1 = str(header_row[1] or "").strip().lower()
        return col1.startswith("prox")

    def load_xl_files(self,lot_number, treeview):
        workbook = openpyxl.load_workbook(I_O.path_1+lot_number+".xlsx")
        sheet= workbook.active
        list_values = list(sheet.values)
        lk=list_values[0] if list_values else ()
        I_O.cone_flip = bool(I_O.cone_flip_from_excel_headers(lk))
        self.clear_treeview(treeview)
       # for col_name in list_values[0]:
       #     treeview.heading(col_name, text=col_name)
        self.setup_tree()
        for item in treeview.get_children():
                treeview.delete(item)            
        for value_tuple in list_values[1:]:
            treeview.insert('', tk.END, values=value_tuple)
        # Keep diagram / instruction labels aligned with the loaded lot orientation.
        if hasattr(self, "setup_pic_Labels"):
            self.setup_pic_Labels()
            if hasattr(self, "clear_params"):
                self.clear_params()

       
       
    def load_lot_number(self):
        directory = PROJECT_PATH / DATA_DIR_NAME
            # Get all filenames with .xlsx extension in the specified directory
        File_Names = [filename for filename in os.listdir(directory) if filename.endswith('.xlsx')]
        
        # Stripping file extensions
        self.entry1_['value'] = [os.path.splitext(filename)[0] for filename in File_Names]




    def write_xl_file(self, lot_num):
        # Persist current Prox Left orientation in the Excel header row.
        if hasattr(self, "setup_tree"):
            self.setup_tree()

        workbook = openpyxl.Workbook()
        sheet= workbook.active
        header_text=[]
        for column in self.tree["columns"]:
            header_text.append(self.tree.heading(column)["text"])
        sheet.append(header_text)   
        for child_item in self.tree.get_children():
            row_values = list(self.tree.item(child_item)["values"])
            i=0
            for value in row_values:
                try: 
                    row_values[i] = float(row_values[i])
                except:
                    row_values[i] =""
                i+=1
            sheet.append(row_values)

        I_O._format_lot_worksheet(sheet)

        if lot_num == "":
            fname= 'Temp_'+ datetime.datetime.now().strftime("%Y-%m-%d_%H_%M")
            workbook.save(I_O.path_1+fname+".xlsx")
        else:
            file_path = I_O.path_1+str(lot_num)+".xlsx"
            tmp_file_path = file_path
            if(os.path.exists(tmp_file_path)):
                n=1
                tmp_file_path = I_O.path_1+str(lot_num)+".bak"+str(n)
                while(os.path.exists(tmp_file_path)):
                    n=n+1
                    tmp_file_path = I_O.path_1+str(lot_num)+".bak"+str(n)
                os.rename(file_path, tmp_file_path)        
            workbook.save(file_path)
        I_O.load_lot_number(self)

    @staticmethod
    def _format_lot_worksheet(sheet):
        """Center cell contents and size columns to fit header/data text."""
        if sheet.max_row < 1 or sheet.max_column < 1:
            return

        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.iter_rows(
            min_row=1,
            max_row=sheet.max_row,
            min_col=1,
            max_col=sheet.max_column,
        ):
            for cell in row:
                cell.alignment = center

        # Size each used column from the widest header or data string in that column.
        for col_idx in range(1, sheet.max_column + 1):
            max_len = 0
            for row_idx in range(1, sheet.max_row + 1):
                value = sheet.cell(row=row_idx, column=col_idx).value
                if value is None:
                    continue
                max_len = max(max_len, len(str(value)))
            # openpyxl width is approximate character units; pad slightly for readability.
            sheet.column_dimensions[get_column_letter(col_idx)].width = max(max_len + 2, 6)

        # Give the header row a bit more height when labels wrap.
        sheet.row_dimensions[1].height = 30
        for row_idx in range(2, sheet.max_row + 1):
            sheet.row_dimensions[row_idx].height = 18
           


class json_cls:

    def __init__(self,master):
      self.m = master

    def write_to_json(
        self,
        bal_pres,
        chuck_pres,
        bal_u,
        clamp_u,
        dia_u,
        pos_u,
        tree_file_name,
        usb_transfer_enabled=None,
        usb_transfer_connected=None,
        wifi_transfer_enabled=None,
        wifi_transfer_connected=None,
    ):
        # Preserve existing transfer flags when callers omit them (e.g. Setup Exit).
        existing = {}
        try:
            with open(str(PROJECT_PATH / STARTUP_JSON_NAME), "r") as f:
                existing = json.load(f)
        except (OSError, json.JSONDecodeError):
            existing = {}

        if usb_transfer_enabled is None:
            usb_transfer_enabled = bool(
                existing.get("usb_transfer_enabled", USB_TRANSFER_ENABLED_DEFAULT)
            )
        if usb_transfer_connected is None:
            usb_transfer_connected = bool(
                existing.get("usb_transfer_connected", USB_TRANSFER_CONNECTED_DEFAULT)
            )
        if wifi_transfer_enabled is None:
            wifi_transfer_enabled = bool(
                existing.get("wifi_transfer_enabled", WIFI_TRANSFER_ENABLED_DEFAULT)
            )
        if wifi_transfer_connected is None:
            wifi_transfer_connected = bool(
                existing.get("wifi_transfer_connected", WIFI_TRANSFER_CONNECTED_DEFAULT)
            )

        data = {
            "target_chuck_pressure": chuck_pres,
            "target_balloon_pressure": bal_pres,
            "cur_bal_units": bal_u,
            "cur_clamp_units": clamp_u,
            "cur_dia_units": dia_u,
            "cur_pos_units": pos_u,
            "cone_flip": I_O.cone_flip,
            "tree_file_name": tree_file_name,
            "usb_transfer_enabled": bool(usb_transfer_enabled),
            "usb_transfer_connected": bool(usb_transfer_connected),
            "wifi_transfer_enabled": bool(wifi_transfer_enabled),
            "wifi_transfer_connected": bool(wifi_transfer_connected),
        }
        with open(str(PROJECT_PATH / STARTUP_JSON_NAME), "w") as f:
            json.dump(data, f, indent=4)

    def read_from_json(self):
        """
        Read data from a JSON file.

        Returns:
        dict: The data read from the file.
        """
        with open(str(PROJECT_PATH / STARTUP_JSON_NAME), "r") as f:
            data = json.load(f)
            I_O.cone_flip = bool(data.get("cone_flip", False))
        data.setdefault("usb_transfer_enabled", USB_TRANSFER_ENABLED_DEFAULT)
        data.setdefault("usb_transfer_connected", USB_TRANSFER_CONNECTED_DEFAULT)
        data.setdefault("wifi_transfer_enabled", WIFI_TRANSFER_ENABLED_DEFAULT)
        data.setdefault("wifi_transfer_connected", WIFI_TRANSFER_CONNECTED_DEFAULT)
        return data

